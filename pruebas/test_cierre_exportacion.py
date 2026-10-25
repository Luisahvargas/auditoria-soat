"""Pruebas del cierre de auditoría y de las exportaciones (Excel y PDF) con datos FICTICIOS."""

import io

import pytest
from openpyxl import load_workbook

from auditoria import catalogos
from auditoria.auditorias import cierre
from auditoria.auditorias import servicio as servicio_auditorias
from auditoria.documentos import inventario, lector_pdf
from auditoria.exportacion.excel import generar_excel
from auditoria.exportacion.informe_pdf import generar_pdf
from auditoria.exportacion.resumen import construir_resumen
from auditoria.extraccion import ficha
from auditoria.hallazgos import servicio as servicio_hallazgos
from auditoria.reglas.motor import ejecutar_reglas
from pruebas import generadores


def _preparar(sesion, config, ids):
    usuario_id, auditoria_id = ids
    inventario.cargar_archivo(
        sesion, config, auditoria_id, usuario_id, "cuenta.pdf", generadores.cuenta_ficticia(lado_informe="izquierda", placa_poliza="XYZ987")
    )
    ejecutar_reglas(sesion, auditoria_id, usuario_id)
    sesion.commit()


def _revisar_todo(sesion, ids):
    usuario_id, auditoria_id = ids
    for h in servicio_hallazgos.pendientes_de_revision(sesion, auditoria_id):
        servicio_hallazgos.revisar(sesion, h.id, usuario_id, catalogos.REVISION_CONFIRMADO)
    sesion.commit()


def _faltantes(sesion, auditoria_id):
    return {r.nombre for r in cierre.requisitos_de_cierre(sesion, auditoria_id) if r.obligatorio and not r.cumplido}


# ---------------- Cierre ----------------


def test_no_se_finaliza_con_pasos_obligatorios_pendientes(sesion, config_temporal, usuario_y_auditoria):
    usuario_id, auditoria_id = usuario_y_auditoria
    assert "Documentos cargados" in _faltantes(sesion, auditoria_id)

    _preparar(sesion, config_temporal, usuario_y_auditoria)
    faltan = _faltantes(sesion, auditoria_id)
    assert faltan == {"Hallazgos revisados por el auditor", "Conclusión del auditor registrada"}
    with pytest.raises(servicio_auditorias.ErrorAuditoria, match="No se puede finalizar"):
        cierre.finalizar(sesion, auditoria_id, usuario_id)


def test_analisis_desactualizado_tras_nueva_carga(sesion, config_temporal, usuario_y_auditoria):
    from datetime import timedelta

    usuario_id, auditoria_id = usuario_y_auditoria
    _preparar(sesion, config_temporal, usuario_y_auditoria)
    assert "Análisis automático actualizado" not in _faltantes(sesion, auditoria_id)

    inventario.cargar_archivo(sesion, config_temporal, auditoria_id, usuario_id, "otro.pdf", generadores.pdf_con_texto())
    nuevo = inventario.listar_documentos(sesion, auditoria_id)[-1]
    nuevo.fecha_carga += timedelta(minutes=5)  # Simula que se cargó después del último análisis
    sesion.commit()
    assert "Análisis automático actualizado" in _faltantes(sesion, auditoria_id)


def test_flujo_completo_finaliza_y_bloquea_cambios(sesion, config_temporal, usuario_y_auditoria):
    usuario_id, auditoria_id = usuario_y_auditoria
    _preparar(sesion, config_temporal, usuario_y_auditoria)
    _revisar_todo(sesion, usuario_y_auditoria)
    servicio_auditorias.guardar_conclusion(sesion, auditoria_id, usuario_id, "Conclusión ficticia del auditor.", "")
    cierre.finalizar(sesion, auditoria_id, usuario_id)
    sesion.commit()

    auditoria = servicio_auditorias.obtener_auditoria(sesion, auditoria_id)
    assert auditoria.estado == catalogos.AUDITORIA_FINALIZADA and auditoria.usuario_finalizacion_id == usuario_id
    with pytest.raises(servicio_auditorias.ErrorAuditoria):
        inventario.cargar_archivo(sesion, config_temporal, auditoria_id, usuario_id, "x.pdf", generadores.pdf_con_texto())
    with pytest.raises(servicio_auditorias.ErrorAuditoria):
        ficha.extraer_auditoria(sesion, auditoria_id, usuario_id)
    with pytest.raises(servicio_auditorias.ErrorAuditoria):
        ejecutar_reglas(sesion, auditoria_id, usuario_id)

    with pytest.raises(servicio_auditorias.ErrorAuditoria, match="motivo"):
        cierre.reabrir(sesion, auditoria_id, usuario_id, "  ")
    cierre.reabrir(sesion, auditoria_id, usuario_id, "Llegó un soporte nuevo (ficticio)")
    sesion.commit()
    assert servicio_auditorias.obtener_auditoria(sesion, auditoria_id).estado == catalogos.AUDITORIA_ABIERTA


# ---------------- Exportaciones ----------------


def test_excel_tiene_hojas_y_coincide_con_el_resumen(sesion, config_temporal, usuario_y_auditoria):
    _, auditoria_id = usuario_y_auditoria
    _preparar(sesion, config_temporal, usuario_y_auditoria)
    resumen = construir_resumen(sesion, auditoria_id)
    libro = load_workbook(io.BytesIO(generar_excel(resumen)))

    assert libro.sheetnames == ["Resumen", "Hallazgos", "Evidencias", "Ficha", "Inventario"]
    hallazgos = libro["Hallazgos"]
    assert hallazgos.max_row - 1 == len(resumen.hallazgos)
    ids = {fila[0].value for fila in hallazgos.iter_rows(min_row=2)}
    assert ids == {h.id for h in resumen.hallazgos}
    origenes = {fila[1].value for fila in hallazgos.iter_rows(min_row=2)}
    assert catalogos.ORIGEN_REGLA in origenes
    assert libro["Inventario"]["A2"].value == "cuenta.pdf"
    valores_resumen = {fila[0].value: fila[1].value for fila in libro["Resumen"].iter_rows()}
    assert valores_resumen["Hallazgos que requieren atención"] == len(resumen.hallazgos_atencion)
    assert "BORRADOR" in valores_resumen["Estado"]
    evidencias = libro["Evidencias"]
    assert {fila[0].value for fila in evidencias.iter_rows(min_row=2)} == {"cuenta.pdf"}
    assert all(isinstance(fila[1].value, int) for fila in evidencias.iter_rows(min_row=2))


def test_excel_no_convierte_texto_en_formula(sesion, config_temporal, usuario_y_auditoria):
    usuario_id, auditoria_id = usuario_y_auditoria
    _preparar(sesion, config_temporal, usuario_y_auditoria)
    servicio_hallazgos.agregar_hallazgo_auditor(
        sesion, auditoria_id, usuario_id, "=HYPERLINK(\"http://x\")", "Otra", "Baja", catalogos.POSIBLE_INCONSISTENCIA,
        "Texto ficticio", [{"archivo": "cuenta.pdf", "documento_id": None, "pagina": 1, "texto": "=1+1"}],
    )
    sesion.commit()
    libro = load_workbook(io.BytesIO(generar_excel(construir_resumen(sesion, auditoria_id))))
    celdas = [c for fila in libro.worksheets for f in fila.iter_rows() for c in f if isinstance(c.value, str) and c.value.startswith("=")]
    assert celdas and all(c.data_type == "s" for c in celdas)


def test_pdf_se_abre_y_diferencia_los_tres_origenes(sesion, config_temporal, usuario_y_auditoria, tmp_path):
    usuario_id, auditoria_id = usuario_y_auditoria
    _preparar(sesion, config_temporal, usuario_y_auditoria)
    servicio_auditorias.guardar_conclusion(sesion, auditoria_id, usuario_id, "Conclusión <ficticia> & con símbolos «especiales»", "")
    sesion.commit()
    resumen = construir_resumen(sesion, auditoria_id)

    ruta = tmp_path / "informe.pdf"
    ruta.write_bytes(generar_pdf(resumen))
    texto = " ".join(" ".join(p.texto.split()) for p in lector_pdf.leer_paginas(ruta))

    assert lector_pdf.contar_paginas(ruta) >= 2
    assert "PRUEBA-001" in texto and "BORRADOR" in texto
    assert "Conclusión <ficticia> & con símbolos «especiales»" in texto
    for seccion in ("Resultados de reglas automáticas", "Sugerencias de IA", "Hallazgos registrados por el auditor"):
        assert seccion in texto
    assert "No constituye glosa" in texto
    assert "cuenta.pdf" in texto
    assert "glosa definitiva" not in texto.lower() and "rechazada" not in texto.lower()


def test_resumen_indica_que_no_hubo_ia(sesion, config_temporal, usuario_y_auditoria):
    _preparar(sesion, config_temporal, usuario_y_auditoria)
    resumen = construir_resumen(sesion, usuario_y_auditoria[1])
    assert not resumen.hubo_ia
    libro = load_workbook(io.BytesIO(generar_excel(resumen)))
    valores = {fila[0].value: fila[1].value for fila in libro["Resumen"].iter_rows()}
    assert valores["Análisis de IA"] == "No se ejecutó análisis de IA"
