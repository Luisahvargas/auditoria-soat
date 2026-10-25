"""Exportación a Excel: resumen, hallazgos detallados, evidencias, ficha e inventario.

El texto que viene de los documentos se escribe siempre como texto (nunca como fórmula).
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from auditoria import catalogos
from auditoria.exportacion.resumen import ResumenAuditoria, lineas_evidencia

ENCABEZADO = PatternFill("solid", fgColor="1F4E78")
FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
AVISO = (
    "Documento de apoyo a la auditoría. No constituye glosa ni rechazo de la cuenta: "
    "la decisión corresponde al auditor."
)


def _escribir(hoja, fila: int, columna: int, valor):
    celda = hoja.cell(row=fila, column=columna, value=valor)
    if isinstance(valor, str):
        celda.data_type = "s"  # Evita que un texto que empiece por "=" se interprete como fórmula
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    return celda


def _tabla(hoja, encabezados: list[str], filas: list[list], anchos: list[int], fila_inicio: int = 1) -> None:
    for columna, titulo in enumerate(encabezados, start=1):
        celda = _escribir(hoja, fila_inicio, columna, titulo)
        celda.fill, celda.font = ENCABEZADO, FUENTE_ENCABEZADO
    for i, fila in enumerate(filas, start=fila_inicio + 1):
        for columna, valor in enumerate(fila, start=1):
            _escribir(hoja, i, columna, valor)
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    hoja.freeze_panes = hoja.cell(row=fila_inicio + 1, column=1)
    if filas:
        hoja.auto_filter.ref = f"A{fila_inicio}:{get_column_letter(len(encabezados))}{fila_inicio + len(filas)}"


def generar_excel(resumen: ResumenAuditoria) -> bytes:
    libro = Workbook()

    # --- Resumen ---
    hoja = libro.active
    hoja.title = "Resumen"
    estado = resumen.estado if resumen.finalizada else f"{resumen.estado} (BORRADOR: auditoría no finalizada)"
    filas = [
        ("Auditoría", resumen.identificador),
        ("Descripción", resumen.descripcion),
        ("Estado", estado),
        ("Perfil de aseguradora", resumen.perfil_aseguradora),
        ("Creada por", f"{resumen.creada_por} el {resumen.fecha_creacion:%Y-%m-%d %H:%M}"),
        ("Finalizada por", f"{resumen.finalizada_por} el {resumen.fecha_finalizacion:%Y-%m-%d %H:%M}" if resumen.fecha_finalizacion else ""),
        ("Generado", f"{resumen.generado:%Y-%m-%d %H:%M}"),
        ("Aviso", AVISO),
        ("", ""),
        ("Archivos", len(resumen.documentos)),
        ("Páginas", sum(d.paginas or 0 for d in resumen.documentos)),
        ("Páginas sin texto (no revisadas por reglas)", sum(d.sin_texto for d in resumen.documentos)),
        ("Edad a la fecha del accidente", resumen.edad_al_accidente),
        ("Último análisis automático", resumen.ultima_ejecucion),
        ("Reglas ejecutadas", resumen.reglas_ejecutadas),
        ("Análisis de IA", "Sí (sugerencias pendientes de revisión humana)" if resumen.hubo_ia else "No se ejecutó análisis de IA"),
        ("Tokens de API de IA (entrada / salida)", f"{resumen.tokens_ia[0]} / {resumen.tokens_ia[1]}"),
        ("", ""),
        ("Hallazgos que requieren atención", len(resumen.hallazgos_atencion)),
        (f"  · {catalogos.ORIGEN_REGLA}", resumen.contar(catalogos.ORIGEN_REGLA)),
        (f"  · {catalogos.ORIGEN_IA}", resumen.contar(catalogos.ORIGEN_IA)),
        (f"  · {catalogos.ORIGEN_AUDITOR}", resumen.contar(catalogos.ORIGEN_AUDITOR)),
        ("  · Pendientes de revisión", resumen.contar(revision=catalogos.REVISION_PENDIENTE)),
        ("  · Confirmados", resumen.contar(revision=catalogos.REVISION_CONFIRMADO)),
        ("  · Descartados", resumen.contar(revision=catalogos.REVISION_DESCARTADO)),
        ("  · Modificados", resumen.contar(revision=catalogos.REVISION_MODIFICADO)),
        ("", ""),
        ("Conclusión del auditor", resumen.conclusion_auditor),
        ("Observaciones del auditor", resumen.observaciones_auditor),
    ]
    for i, (etiqueta, valor) in enumerate(filas, start=1):
        _escribir(hoja, i, 1, etiqueta).font = Font(bold=True)
        _escribir(hoja, i, 2, valor)
    hoja.column_dimensions["A"].width = 45
    hoja.column_dimensions["B"].width = 100

    # --- Hallazgos ---
    hoja = libro.create_sheet("Hallazgos")
    _tabla(
        hoja,
        ["ID", "Origen", "Regla", "Título", "Categoría", "Severidad", "Estado (resultado)", "Estado final", "Requiere atención",
         "Revisión", "Revisor", "Fecha revisión", "Comentario del auditor", "Explicación", "Referencia normativa", "Evidencias (archivo, página)"],
        [
            [h.id, h.origen, h.regla, h.titulo, h.categoria, h.severidad, h.estado_resultado, h.estado_final,
             "Sí" if h.requiere_atencion else "No", h.revision, h.revisor,
             h.fecha_revision.strftime("%Y-%m-%d %H:%M") if h.fecha_revision else "", h.comentario, h.explicacion,
             h.referencia_normativa, "; ".join(f"{e.get('archivo')} p. {e.get('pagina') or '-'}" for e in h.evidencias)]
            for h in resumen.hallazgos
        ],
        [6, 20, 10, 45, 22, 10, 28, 28, 10, 26, 18, 16, 40, 70, 22, 40],
    )

    # --- Evidencias ---
    hoja = libro.create_sheet("Evidencias")
    _tabla(
        hoja,
        ["Archivo", "Página del archivo", "ID hallazgo", "Origen", "Estado", "Texto de la evidencia"],
        [[f["archivo"], f["pagina"], f["hallazgo"], f["origen"], f["estado"], f["texto"]] for f in lineas_evidencia(resumen)],
        [30, 12, 10, 20, 28, 90],
    )

    # --- Ficha ---
    hoja = libro.create_sheet("Ficha")
    _tabla(
        hoja,
        ["Bloque", "Campo", "Valor(es)", "Fuentes", "Valores distintos", "Corregidos por el auditor"],
        [[f.bloque, f.campo, f.valores, f.fuentes, "Sí" if f.discrepancia else "", f.corregidos or ""] for f in resumen.ficha],
        [26, 40, 60, 10, 16, 14],
    )

    # --- Inventario ---
    hoja = libro.create_sheet("Inventario")
    _tabla(
        hoja,
        ["Archivo", "Tipo", "Estado", "Páginas", "Sin texto", "Sin confirmar clasificación", "Ilegibles", "SHA-256", "Fecha de carga", "Detalle"],
        [
            [d.archivo, d.tipo, d.estado, d.paginas, d.sin_texto, d.sin_confirmar, d.ilegibles, d.hash_sha256,
             d.fecha_carga.strftime("%Y-%m-%d %H:%M"), d.mensaje]
            for d in resumen.documentos
        ],
        [35, 8, 20, 10, 10, 14, 10, 66, 17, 60],
    )

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
